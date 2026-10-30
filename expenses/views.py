import datetime
from decimal import Decimal
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.core.paginator import Paginator
from django.utils import timezone
from django.db.models import Sum, Q, Count

from .models import (
    CashEmployee, OpeningBalance, ExpenseSite, ExpenseCategory,
    CashTransaction, EmployeeTransfer, ExpenseAuditLog
)
from .services import AccountingService


@login_required
def expense_dashboard_view(request):
    """Executive Dashboard for Employee Cash & Expense Management."""
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')
    
    start_date = None
    end_date = None
    if start_date_str:
        try:
            start_date = datetime.datetime.strptime(start_date_str, '%Y-%m-%d').date()
        except ValueError:
            pass
    if end_date_str:
        try:
            end_date = datetime.datetime.strptime(end_date_str, '%Y-%m-%d').date()
        except ValueError:
            pass

    summary = AccountingService.get_dashboard_summary(start_date=start_date, end_date=end_date)
    
    # Recent 15 transactions
    recent_txns = (
        CashTransaction.objects.select_related('employee', 'site', 'category')
        .order_by('-transaction_date', '-id')[:15]
    )

    context = {
        'summary': summary,
        'recent_txns': recent_txns,
        'start_date': start_date_str or '',
        'end_date': end_date_str or '',
        'page_title': 'Employee Cash & Expense Dashboard',
    }
    return render(request, 'expenses/dashboard.html', context)


@login_required
def ledger_list_view(request):
    """Lists all employee cash accounts and their balances."""
    employees = CashEmployee.objects.filter(is_active=True).prefetch_related('opening_balances')
    emp_cards = []
    
    for emp in employees:
        op = emp.opening_balances.first()
        op_amount = op.amount if op else Decimal('0.00')

        c_in = emp.transactions.filter(
            txn_type__in=[CashTransaction.TxnType.CASH_IN, CashTransaction.TxnType.TRANSFER_IN]
        ).aggregate(tot=Sum('amount'))['tot'] or Decimal('0.00')

        exp = emp.transactions.filter(
            txn_type=CashTransaction.TxnType.EXPENSE
        ).aggregate(tot=Sum('amount'))['tot'] or Decimal('0.00')

        c_out = emp.transactions.filter(
            txn_type__in=[CashTransaction.TxnType.CASH_OUT, CashTransaction.TxnType.TRANSFER_OUT]
        ).aggregate(tot=Sum('amount'))['tot'] or Decimal('0.00')

        emp_cards.append({
            'employee': emp,
            'opening_balance': op_amount,
            'cash_in': c_in,
            'expenses': exp,
            'cash_out': c_out,
            'current_balance': emp.cached_balance,
            'txn_count': emp.transactions.count(),
        })

    context = {
        'emp_cards': emp_cards,
        'page_title': 'Employee Cash Ledgers',
    }
    return render(request, 'expenses/ledger_list.html', context)


@login_required
def employee_ledger_view(request, emp_id):
    """Detailed transaction ledger for a single employee."""
    employee = get_object_or_404(CashEmployee, id=emp_id)

    # Search & Date Filters
    query = request.GET.get('q', '').strip()
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')
    txn_type_filter = request.GET.get('txn_type')
    site_filter = request.GET.get('site')
    category_filter = request.GET.get('category')

    txns_qs = employee.transactions.select_related('site', 'category').order_by('transaction_date', 'id')

    if query:
        txns_qs = txns_qs.filter(
            Q(particulars__icontains=query) |
            Q(cash_in_source__icontains=query) |
            Q(cash_out_recipient__icontains=query) |
            Q(voucher_no__icontains=query) |
            Q(site__name__icontains=query) |
            Q(category__name__icontains=query)
        )

    if start_date_str:
        try:
            s_date = datetime.datetime.strptime(start_date_str, '%Y-%m-%d').date()
            txns_qs = txns_qs.filter(transaction_date__gte=s_date)
        except ValueError:
            pass

    if end_date_str:
        try:
            e_date = datetime.datetime.strptime(end_date_str, '%Y-%m-%d').date()
            txns_qs = txns_qs.filter(transaction_date__lte=e_date)
        except ValueError:
            pass

    if txn_type_filter:
        txns_qs = txns_qs.filter(txn_type=txn_type_filter)

    if site_filter:
        txns_qs = txns_qs.filter(site_id=site_filter)

    if category_filter:
        txns_qs = txns_qs.filter(category_id=category_filter)

    # Filter by Expense Head (Labour, Machinery, Other)
    head_filter = request.GET.get('head', '').strip().upper()
    if head_filter == 'LABOUR':
        txns_qs = txns_qs.filter(
            Q(category__group_type=ExpenseCategory.GroupType.LABOUR) |
            Q(particulars__icontains='labour') |
            Q(particulars__icontains='majdoor')
        )
    elif head_filter == 'MACHINERY':
        txns_qs = txns_qs.filter(
            Q(category__group_type__in=[ExpenseCategory.GroupType.MACHINERY, ExpenseCategory.GroupType.VEHICLE]) |
            Q(particulars__icontains='machin') |
            Q(particulars__icontains='diesel') |
            Q(particulars__icontains='jcb')
        ).exclude(
            Q(category__group_type=ExpenseCategory.GroupType.LABOUR) |
            Q(particulars__icontains='labour') |
            Q(particulars__icontains='majdoor')
        )
    elif head_filter == 'OTHER':
        txns_qs = txns_qs.filter(txn_type=CashTransaction.TxnType.EXPENSE).exclude(
            Q(category__group_type__in=[ExpenseCategory.GroupType.LABOUR, ExpenseCategory.GroupType.MACHINERY, ExpenseCategory.GroupType.VEHICLE]) |
            Q(particulars__icontains='labour') |
            Q(particulars__icontains='majdoor') |
            Q(particulars__icontains='machin') |
            Q(particulars__icontains='diesel') |
            Q(particulars__icontains='jcb')
        )

    # High-level expenditure breakdown: Labour, Machinery, Other
    breakdown = employee.get_expense_breakdown(start_date=start_date_str, end_date=end_date_str)

    # Calculate summary metrics for this employee
    op = employee.opening_balances.first()
    opening_bal = op.amount if op else Decimal('0.00')

    total_in = employee.transactions.filter(
        txn_type__in=[CashTransaction.TxnType.CASH_IN, CashTransaction.TxnType.TRANSFER_IN]
    ).aggregate(tot=Sum('amount'))['tot'] or Decimal('0.00')

    total_exp = employee.transactions.filter(
        txn_type=CashTransaction.TxnType.EXPENSE
    ).aggregate(tot=Sum('amount'))['tot'] or Decimal('0.00')

    total_out = employee.transactions.filter(
        txn_type__in=[CashTransaction.TxnType.CASH_OUT, CashTransaction.TxnType.TRANSFER_OUT]
    ).aggregate(tot=Sum('amount'))['tot'] or Decimal('0.00')

    # Pagination: 100 per page for high performance
    paginator = Paginator(txns_qs, 100)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    # Form master options for modals
    sites = ExpenseSite.objects.filter(is_active=True)
    categories = ExpenseCategory.objects.filter(is_active=True)
    other_employees = CashEmployee.objects.filter(is_active=True).exclude(id=employee.id)

    context = {
        'employee': employee,
        'page_obj': page_obj,
        'opening_balance': opening_bal,
        'total_in': total_in,
        'total_exp': total_exp,
        'total_out': total_out,
        'current_balance': employee.cached_balance,
        'breakdown': breakdown,
        'selected_head': head_filter,
        'sites': sites,
        'categories': categories,
        'other_employees': other_employees,
        'q': query,
        'start_date': start_date_str or '',
        'end_date': end_date_str or '',
        'txn_type_filter': txn_type_filter or '',
        'site_filter': site_filter or '',
        'category_filter': category_filter or '',
        'page_title': f"{employee.name}'s Cash Ledger",
    }
    return render(request, 'expenses/employee_ledger.html', context)


@login_required
def add_cash_in_view(request, emp_id):
    """Handles adding Cash In entry to an employee ledger."""
    if request.method != 'POST':
        return redirect('expenses:employee_ledger', emp_id=emp_id)

    employee = get_object_or_404(CashEmployee, id=emp_id)
    amount_str = request.POST.get('amount')
    date_str = request.POST.get('transaction_date')
    source = request.POST.get('cash_in_source', '').strip()
    remarks = request.POST.get('remarks', '').strip()

    try:
        amount = Decimal(amount_str)
        if amount <= 0:
            raise ValueError("Amount must be greater than zero.")
        txn_date = datetime.datetime.strptime(date_str, '%Y-%m-%d').date()
    except (ValueError, TypeError) as e:
        messages.error(request, f"Invalid Cash In entry: {e}")
        return redirect('expenses:employee_ledger', emp_id=emp_id)

    AccountingService.create_cash_in(
        employee=employee, amount=amount, date=txn_date,
        source=source, remarks=remarks, user=request.user
    )
    messages.success(request, f"Successfully recorded Cash In of ₹{amount:,.2f} for {employee.name}.")
    return redirect('expenses:employee_ledger', emp_id=emp_id)


@login_required
def add_expense_view(request, emp_id):
    """Handles adding an Expense entry to an employee ledger."""
    if request.method != 'POST':
        return redirect('expenses:employee_ledger', emp_id=emp_id)

    employee = get_object_or_404(CashEmployee, id=emp_id)
    amount_str = request.POST.get('amount')
    date_str = request.POST.get('transaction_date')
    site_id = request.POST.get('site')
    category_id = request.POST.get('category')
    particulars = request.POST.get('particulars', '').strip()
    voucher_no = request.POST.get('voucher_no', '').strip()
    remarks = request.POST.get('remarks', '').strip()

    try:
        amount = Decimal(amount_str)
        if amount <= 0:
            raise ValueError("Expense amount must be greater than zero.")
        txn_date = datetime.datetime.strptime(date_str, '%Y-%m-%d').date()
    except (ValueError, TypeError) as e:
        messages.error(request, f"Invalid Expense entry: {e}")
        return redirect('expenses:employee_ledger', emp_id=emp_id)

    site_obj = ExpenseSite.objects.filter(id=site_id).first() if site_id else None
    cat_obj = ExpenseCategory.objects.filter(id=category_id).first() if category_id else None
    given_to_id = request.POST.get('given_to_employee')
    given_to_employee = CashEmployee.objects.filter(id=given_to_id, is_active=True).first() if given_to_id else None

    AccountingService.create_expense(
        employee=employee, amount=amount, date=txn_date,
        site=site_obj, category=cat_obj, particulars=particulars,
        voucher_no=voucher_no, remarks=remarks,
        given_to_employee=given_to_employee, user=request.user
    )
    if given_to_employee:
        messages.success(
            request,
            f"Successfully recorded ₹{amount:,.2f} given to {given_to_employee.name}. "
            f"An incoming entry has been automatically credited to {given_to_employee.name}'s ledger on {txn_date}."
        )
    else:
        messages.success(request, f"Successfully recorded Expense of ₹{amount:,.2f} for {employee.name}.")
    return redirect('expenses:employee_ledger', emp_id=emp_id)


@login_required
def add_cash_out_view(request, emp_id):
    """Handles adding Cash Out payment to an employee ledger."""
    if request.method != 'POST':
        return redirect('expenses:employee_ledger', emp_id=emp_id)

    employee = get_object_or_404(CashEmployee, id=emp_id)
    amount_str = request.POST.get('amount')
    date_str = request.POST.get('transaction_date')
    recipient = request.POST.get('cash_out_recipient', '').strip()
    recipient_emp_id = request.POST.get('recipient_employee')
    particulars = request.POST.get('particulars', '').strip()
    remarks = request.POST.get('remarks', '').strip()

    try:
        amount = Decimal(amount_str)
        if amount <= 0:
            raise ValueError("Amount must be greater than zero.")
        txn_date = datetime.datetime.strptime(date_str, '%Y-%m-%d').date()
    except (ValueError, TypeError) as e:
        messages.error(request, f"Invalid Cash Out entry: {e}")
        return redirect('expenses:employee_ledger', emp_id=emp_id)

    recipient_employee = CashEmployee.objects.filter(id=recipient_emp_id, is_active=True).first() if recipient_emp_id else None

    AccountingService.create_cash_out(
        employee=employee, amount=amount, date=txn_date,
        recipient=recipient, particulars=particulars, remarks=remarks,
        recipient_employee=recipient_employee, user=request.user
    )
    if recipient_employee:
        messages.success(
            request,
            f"Successfully recorded ₹{amount:,.2f} given to {recipient_employee.name}. "
            f"An incoming entry has been automatically credited to {recipient_employee.name}'s ledger on {txn_date}."
        )
    else:
        messages.success(request, f"Successfully recorded Cash Out of ₹{amount:,.2f} to {recipient}.")
    return redirect('expenses:employee_ledger', emp_id=emp_id)


@login_required
def add_transfer_view(request):
    """Handles internal employee-to-employee transfer atomically."""
    if request.method != 'POST':
        return redirect('expenses:transfers_list')

    sender_id = request.POST.get('sender_id')
    receiver_id = request.POST.get('receiver_id')
    amount_str = request.POST.get('amount')
    date_str = request.POST.get('transfer_date')
    payment_mode = request.POST.get('payment_mode', 'Paytm').strip()
    reference_no = request.POST.get('reference_no', '').strip()
    purpose = request.POST.get('purpose', '').strip()
    remarks = request.POST.get('remarks', '').strip()
    full_remarks = f"[{purpose}] {remarks}".strip() if purpose else remarks
    next_url = request.POST.get('next_url')

    sender = get_object_or_404(CashEmployee, id=sender_id)
    receiver = get_object_or_404(CashEmployee, id=receiver_id)

    try:
        amount = Decimal(amount_str)
        if amount <= 0:
            raise ValueError("Transfer amount must be greater than zero.")
        t_date = datetime.datetime.strptime(date_str, '%Y-%m-%d').date() if date_str else timezone.localdate()
        
        AccountingService.create_employee_transfer(
            sender=sender, receiver=receiver, amount=amount,
            date=t_date, payment_mode=payment_mode,
            reference_no=reference_no, remarks=full_remarks, user=request.user
        )
        messages.success(request, f"Transferred ₹{amount:,.2f} from {sender.name} to {receiver.name} on {t_date}. Automatically synced in both ledgers.")
    except Exception as e:
        messages.error(request, f"Transfer failed: {e}")

    if next_url:
        return redirect(next_url)
    return redirect('expenses:transfers_list')


@login_required
def delete_transaction_view(request, txn_id):
    """Deletes a transaction and recalculates the ledger."""
    if request.method != 'POST':
        return redirect('expenses:dashboard')

    txn = get_object_or_404(CashTransaction, id=txn_id)
    employee = txn.employee
    emp_id = employee.id
    amt = txn.amount
    t_type = txn.get_txn_type_display()

    # If linked to a transfer, delete both legs
    if hasattr(txn, 'transfer_out_parent') and txn.transfer_out_parent:
        tr = txn.transfer_out_parent
        receiver = tr.receiver
        if tr.credit_transaction:
            tr.credit_transaction.delete()
        tr.delete()
        txn.delete()
        AccountingService.recalculate_employee_ledger(receiver)
    elif hasattr(txn, 'transfer_in_parent') and txn.transfer_in_parent:
        tr = txn.transfer_in_parent
        sender = tr.sender
        if tr.debit_transaction:
            tr.debit_transaction.delete()
        tr.delete()
        txn.delete()
        AccountingService.recalculate_employee_ledger(sender)
    else:
        txn.delete()

    AccountingService.recalculate_employee_ledger(employee)
    messages.warning(request, f"Deleted {t_type} entry of ₹{amt:,.2f}. Ledger recalculated.")
    return redirect('expenses:employee_ledger', emp_id=emp_id)


@login_required
def transfers_list_view(request):
    """Shows all internal transfers between employees."""
    transfers = (
        EmployeeTransfer.objects.select_related('sender', 'receiver')
        .order_by('-transfer_date', '-id')
    )
    employees = CashEmployee.objects.filter(is_active=True)

    paginator = Paginator(transfers, 50)
    page_number = request.GET.get('page', 1)
    page_obj = paginator.get_page(page_number)

    context = {
        'page_obj': page_obj,
        'employees': employees,
        'page_title': 'Employee Transfers History',
    }
    return render(request, 'expenses/transfers_list.html', context)


@login_required
def site_reports_view(request):
    """Site-wise expense analysis and breakdown."""
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')

    q_filter = Q(txn_type=CashTransaction.TxnType.EXPENSE, site__isnull=False)
    if start_date_str:
        q_filter &= Q(transaction_date__gte=start_date_str)
    if end_date_str:
        q_filter &= Q(transaction_date__lte=end_date_str)

    sites_summary = (
        CashTransaction.objects.filter(q_filter)
        .values('site__id', 'site__name', 'site__location')
        .annotate(total_spent=Sum('amount'), txn_count=Count('id'))
        .order_by('-total_spent')
    )

    context = {
        'sites_summary': sites_summary,
        'start_date': start_date_str or '',
        'end_date': end_date_str or '',
        'page_title': 'Site-Wise Expenditure Report',
    }
    return render(request, 'expenses/site_reports.html', context)


@login_required
def category_reports_view(request):
    """Category-wise expense analysis and breakdown."""
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')

    q_filter = Q(txn_type=CashTransaction.TxnType.EXPENSE, category__isnull=False)
    if start_date_str:
        q_filter &= Q(transaction_date__gte=start_date_str)
    if end_date_str:
        q_filter &= Q(transaction_date__lte=end_date_str)

    cat_summary = (
        CashTransaction.objects.filter(q_filter)
        .values('category__id', 'category__name', 'category__group_type')
        .annotate(total_spent=Sum('amount'), txn_count=Count('id'))
        .order_by('-total_spent')
    )

    context = {
        'cat_summary': cat_summary,
        'start_date': start_date_str or '',
        'end_date': end_date_str or '',
        'page_title': 'Category-Wise Expenditure Report',
    }
    return render(request, 'expenses/category_reports.html', context)


@login_required
def export_ledger_excel_view(request, emp_id):
    """Generates a professional Excel workbook for an employee ledger."""
    employee = get_object_or_404(CashEmployee, id=emp_id)
    op = employee.opening_balances.first()
    opening_bal = op.amount if op else Decimal('0.00')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{employee.name} Ledger"

    # Styling definitions
    title_font = Font(name='Calibri', size=15, bold=True, color='1F4E79')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='203764', end_color='203764', fill_type='solid')
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Title Banner
    ws['A1'] = f"VIKAS ENGINEERING - EMPLOYEE CASH LEDGER: {employee.name.upper()}"
    ws['A1'].font = title_font
    ws.merge_cells('A1:J1')

    ws['A2'] = f"Financial Period: 2026-27 | Designation: {employee.designation or 'Operations'} | Opening Balance: INR {opening_bal:,.2f} | Current Balance: INR {employee.cached_balance:,.2f}"
    ws['A2'].font = Font(size=10, italic=True)
    ws.merge_cells('A2:J2')

    headers = [
        "Date", "Cash In Source", "Cash In (INR)", "Site Allocated",
        "Category", "Particulars / Work Details", "Voucher No",
        "Expense (INR)", "Cash Out / To", "Cash Out (INR)", "Running Balance (INR)"
    ]

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    # Row 5: Opening Balance row
    ws.cell(row=5, column=1, value="OPENING BALANCE")
    ws.cell(row=5, column=11, value=float(opening_bal))

    txns = employee.transactions.select_related('site', 'category').order_by('transaction_date', 'id')
    row_num = 6

    for txn in txns:
        ws.cell(row=row_num, column=1, value=txn.transaction_date.strftime('%d-%m-%Y'))
        
        # Cash in
        if txn.txn_type in (CashTransaction.TxnType.CASH_IN, CashTransaction.TxnType.TRANSFER_IN):
            ws.cell(row=row_num, column=2, value=txn.cash_in_source or '')
            ws.cell(row=row_num, column=3, value=float(txn.amount))
        
        # Expense
        if txn.txn_type == CashTransaction.TxnType.EXPENSE:
            ws.cell(row=row_num, column=4, value=txn.site.name if txn.site else '')
            ws.cell(row=row_num, column=5, value=txn.category.name if txn.category else '')
            ws.cell(row=row_num, column=6, value=txn.particulars or '')
            ws.cell(row=row_num, column=7, value=txn.voucher_no or '')
            ws.cell(row=row_num, column=8, value=float(txn.amount))
        elif txn.particulars and txn.txn_type != CashTransaction.TxnType.EXPENSE:
            ws.cell(row=row_num, column=6, value=txn.particulars or '')

        # Cash out
        if txn.txn_type in (CashTransaction.TxnType.CASH_OUT, CashTransaction.TxnType.TRANSFER_OUT):
            ws.cell(row=row_num, column=9, value=txn.cash_out_recipient or '')
            ws.cell(row=row_num, column=10, value=float(txn.amount))

        ws.cell(row=row_num, column=11, value=float(txn.running_balance or 0.0))

        # Borders
        for c in range(1, 12):
            ws.cell(row=row_num, column=c).border = border_thin

        row_num += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"{employee.name}_Ledger_{timezone.localdate()}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@login_required
def export_consolidated_excel_view(request):
    """Exports company-wide consolidated transactions report."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Company Consolidated Expenses"

    headers = [
        "Date", "Employee", "Type", "Amount (INR)", "Site",
        "Category", "Particulars", "Source / Recipient", "Voucher No"
    ]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E79')

    txns = CashTransaction.objects.select_related('employee', 'site', 'category').order_by('transaction_date', 'id')
    for r, t in enumerate(txns, 2):
        ws.cell(r, 1, t.transaction_date.strftime('%d-%m-%Y'))
        ws.cell(r, 2, t.employee.name)
        ws.cell(r, 3, t.get_txn_type_display())
        ws.cell(r, 4, float(t.amount))
        ws.cell(r, 5, t.site.name if t.site else '')
        ws.cell(r, 6, t.category.name if t.category else '')
        ws.cell(r, 7, t.particulars or '')
        ws.cell(r, 8, t.cash_in_source or t.cash_out_recipient or '')
        ws.cell(r, 9, t.voucher_no or '')

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="Company_Expenses_{timezone.localdate()}.xlsx"'
    return response
