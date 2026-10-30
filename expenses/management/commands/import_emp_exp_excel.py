import datetime
from decimal import Decimal
import openpyxl
from django.core.management.base import BaseCommand
from django.utils import timezone
from expenses.models import (
    CashEmployee, OpeningBalance, ExpenseSite, ExpenseCategory,
    CashTransaction, EmployeeTransfer
)
from expenses.services import AccountingService


class Command(BaseCommand):
    help = "Imports historical employee cash ledgers and transactions from EMP EXP 26-27.xlsx"

    def add_arguments(self, parser):
        parser.add_argument('--file', type=str, default='EMP EXP 26-27.xlsx', help='Path to Excel workbook')
        parser.add_argument('--dry-run', action='store_true', help='Simulate without committing to DB')

    def handle(self, *args, **options):
        file_path = options['file']
        dry_run = options['dry_run']

        self.stdout.write(f"Loading workbook '{file_path}' (dry_run={dry_run})...")
        wb = openpyxl.load_workbook(file_path, data_only=True)

        target_sheets = ['Bipin', 'Divyansh', 'Vikash', 'Sonu', 'Rakesh', 'Lalit', 'Rohit VRK']
        
        # Opening balances verified from row 3 & row 5 of each sheet
        opening_balances_map = {
            'Bipin': Decimal('-174186.00'),
            'Divyansh': Decimal('-108438.00'),
            'Vikash': Decimal('1791.00'),
            'Sonu': Decimal('0.00'),
            'Rakesh': Decimal('-75478.00'),
            'Lalit': Decimal('796.00'),
            'Rohit VRK': Decimal('-1422.00'),
        }

        if not dry_run:
            # Clear existing data to allow clean re-runs
            self.stdout.write("Clearing existing expenses records for fresh import...")
            EmployeeTransfer.objects.all().delete()
            CashTransaction.objects.all().delete()
            OpeningBalance.objects.all().delete()

        total_imported = 0

        for sheet_name in target_sheets:
            if sheet_name not in wb.sheetnames:
                self.stdout.write(self.style.WARNING(f"Sheet '{sheet_name}' not found, skipping."))
                continue

            ws = wb[sheet_name]
            self.stdout.write(f"\nProcessing sheet: {sheet_name} ({ws.max_row} rows)...")

            # 1. Employee
            if not dry_run:
                emp, _ = CashEmployee.objects.get_or_create(
                    name=sheet_name,
                    defaults={'designation': 'Site / Project Operations', 'is_active': True}
                )
                # Opening balance
                op_amt = opening_balances_map.get(sheet_name, Decimal('0.00'))
                OpeningBalance.objects.update_or_create(
                    employee=emp,
                    financial_year='2026-27',
                    defaults={'amount': op_amt, 'effective_date': datetime.date(2026, 4, 1)}
                )
            else:
                emp = None

            # 2. Identify column mapping from Row 4
            header_row = [ws.cell(4, c).value for c in range(1, ws.max_column + 1)]
            col_map = {}
            for idx, h in enumerate(header_row, 1):
                if not h: continue
                h_str = str(h).strip().upper()
                if 'DATE' in h_str and 'DATE' not in col_map: col_map['DATE'] = idx
                elif h_str == 'CASH IN': col_map['CASH_IN_SRC'] = idx
                elif h_str == 'AMOUNT' and 'CASH_IN_AMT' not in col_map: col_map['CASH_IN_AMT'] = idx
                elif 'SITE' in h_str: col_map['SITE'] = idx
                elif 'CAT' in h_str: col_map['CATEGORY'] = idx
                elif 'PART' in h_str: col_map['PARTICULARS'] = idx
                elif h_str == 'AMOUNT2': col_map['EXPENSE_AMT'] = idx
                elif 'VOU' in h_str: col_map['VOUCHER'] = idx
                elif h_str == 'CASH OUT': col_map['CASH_OUT_TO'] = idx
                elif h_str == 'AMOUNT3': col_map['CASH_OUT_AMT'] = idx
                elif 'BAL' in h_str: col_map['BALANCE'] = idx

            sheet_count = 0
            last_valid_date = datetime.date(2026, 4, 1)

            # 3. Read rows from 6 onwards
            for r in range(6, ws.max_row + 1):
                # Date parsing
                raw_d = ws.cell(r, col_map.get('DATE', 1)).value
                txn_date = None
                if isinstance(raw_d, datetime.datetime):
                    txn_date = raw_d.date()
                elif isinstance(raw_d, datetime.date):
                    txn_date = raw_d
                elif isinstance(raw_d, str):
                    for fmt in ('%d.%m.%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%Y'):
                        try:
                            txn_date = datetime.datetime.strptime(raw_d.strip(), fmt).date()
                            break
                        except ValueError:
                            pass

                cin_src = ws.cell(r, col_map.get('CASH_IN_SRC', 2)).value
                cin_amt_val = ws.cell(r, col_map.get('CASH_IN_AMT', 3)).value
                site_val = ws.cell(r, col_map.get('SITE', 4)).value
                cat_val = ws.cell(r, col_map.get('CATEGORY', 5)).value
                part_val = ws.cell(r, col_map.get('PARTICULARS', 6)).value
                exp_amt_val = ws.cell(r, col_map.get('EXPENSE_AMT', 7)).value
                vou_val = ws.cell(r, col_map.get('VOUCHER', 8)).value if 'VOUCHER' in col_map else None
                cout_to = ws.cell(r, col_map.get('CASH_OUT_TO', 8 if 'VOUCHER' not in col_map else 9)).value
                cout_amt_val = ws.cell(r, col_map.get('CASH_OUT_AMT', 9 if 'VOUCHER' not in col_map else 10)).value

                def to_decimal(v):
                    if v is None: return Decimal('0.00')
                    if isinstance(v, (int, float)): return Decimal(str(round(v, 2)))
                    try:
                        clean_str = str(v).replace(',', '').strip()
                        return Decimal(clean_str)
                    except:
                        return Decimal('0.00')

                cin_amt = to_decimal(cin_amt_val)
                exp_amt = to_decimal(exp_amt_val)
                cout_amt = to_decimal(cout_amt_val)

                # Skip completely empty rows
                if cin_amt == 0 and exp_amt == 0 and cout_amt == 0 and not any([site_val, cat_val, part_val]):
                    continue

                if not txn_date:
                    txn_date = last_valid_date
                else:
                    last_valid_date = txn_date

                # If Cash In exists
                if cin_amt > 0:
                    src_text = str(cin_src).strip() if cin_src else "Cash Inflow"
                    # Check if incoming transfer
                    is_transfer = any(other.lower() in src_text.lower() for other in target_sheets if other.lower() != sheet_name.lower())
                    txn_t = CashTransaction.TxnType.TRANSFER_IN if is_transfer else CashTransaction.TxnType.CASH_IN
                    
                    if not dry_run:
                        CashTransaction.objects.create(
                            employee=emp,
                            transaction_date=txn_date,
                            txn_type=txn_t,
                            amount=cin_amt,
                            cash_in_source=src_text,
                            particulars=str(part_val).strip() if part_val and exp_amt == 0 else "",
                            excel_sheet_ref=sheet_name,
                            excel_row_ref=r
                        )
                    sheet_count += 1
                    total_imported += 1

                # If Site / Office Expense exists
                if exp_amt > 0:
                    site_obj = AccountingService.get_or_create_site(site_val) if not dry_run else None
                    cat_obj = AccountingService.get_or_create_category(cat_val) if not dry_run else None
                    
                    if not dry_run:
                        CashTransaction.objects.create(
                            employee=emp,
                            transaction_date=txn_date,
                            txn_type=CashTransaction.TxnType.EXPENSE,
                            amount=exp_amt,
                            site=site_obj,
                            category=cat_obj,
                            particulars=str(part_val).strip() if part_val else "",
                            voucher_no=str(vou_val).strip() if vou_val else "",
                            excel_sheet_ref=sheet_name,
                            excel_row_ref=r
                        )
                    sheet_count += 1
                    total_imported += 1

                # If Cash Out exists
                if cout_amt > 0:
                    out_text = str(cout_to).strip() if cout_to else "Cash Out"
                    is_transfer = any(other.lower() in out_text.lower() for other in target_sheets if other.lower() != sheet_name.lower())
                    txn_t = CashTransaction.TxnType.TRANSFER_OUT if is_transfer else CashTransaction.TxnType.CASH_OUT

                    if not dry_run:
                        CashTransaction.objects.create(
                            employee=emp,
                            transaction_date=txn_date,
                            txn_type=txn_t,
                            amount=cout_amt,
                            cash_out_recipient=out_text,
                            particulars=str(part_val).strip() if part_val and exp_amt == 0 else "",
                            excel_sheet_ref=sheet_name,
                            excel_row_ref=r
                        )
                    sheet_count += 1
                    total_imported += 1

            if not dry_run and emp:
                final_bal = AccountingService.recalculate_employee_ledger(emp)
                self.stdout.write(f"  Imported {sheet_count} records for {sheet_name}. Final Calculated Balance: INR {final_bal:,.2f}")
            else:
                self.stdout.write(f"  [Dry Run] Found {sheet_count} records for {sheet_name}")

        self.stdout.write(self.style.SUCCESS(f"\nSuccessfully processed {total_imported} transaction records across 7 sheets!"))
